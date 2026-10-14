from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

FONT = "Times New Roman"
def new_doc(title, subject):
    d = Document()
    for s in d.sections:
        s.left_margin = Inches(1.0); s.right_margin = Inches(1.0)
        s.top_margin = Inches(0.9); s.bottom_margin = Inches(0.9)
    st = d.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
    for lvl, size in ((1, 14), (2, 12), (3, 12)):
        h = d.styles[f"Heading {lvl}"]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0); h.font.italic = False
        h.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
        rf = h.element.rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); h.element.rPr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10); h.paragraph_format.space_after = Pt(6)
    for name in ("List Bullet", "List Number"):
        d.styles[name].font.name = FONT; d.styles[name].font.size = Pt(12)
    cp = d.core_properties
    cp.author = "Vishwak"; cp.last_modified_by = "Vishwak"; cp.title = title; cp.subject = subject
    cp.comments = ""; cp.keywords = "DBMS, PostgreSQL, BOM"; cp.category = "Course Project"
    cp.created = datetime.datetime(2026, 10, 14, 9, 0); cp.modified = datetime.datetime.now()
    return d

def _field(run, instr):
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(it); run._r.append(f2)

def header_footer(d, left_text, start_section=1):
    for i, s in enumerate(d.sections):
        s.different_first_page_header_footer = True
        hp = s.header.paragraphs[0]; hp.text = ""
        r = hp.add_run(left_text); r.font.size = Pt(9); r.font.name = FONT; r.italic = True
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run(); r.font.size = Pt(10); r.font.name = FONT; _field(r, "PAGE")

def para(d, text, bold=False, italic=False, align=None, size=None, after=None):
    p = d.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "j": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def rich(d, parts, align="j"):
    """parts: list of (text, style) where style in '', 'b', 'i', 'c' (code)"""
    p = d.add_paragraph()
    for t, s in parts:
        r = p.add_run(t)
        if "b" in s: r.bold = True
        if "i" in s: r.italic = True
        if "c" in s: r.font.name = "Consolas"; r.font.size = Pt(10.5)
    if align == "j": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def bullet(d, text): 
    p = d.add_paragraph(text, style="List Bullet"); p.paragraph_format.space_after = Pt(3); return p

def table(d, headers, rows, widths=None, size=10.5, mono_cols=()):
    t = d.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "E7E6E6"); c._tc.get_or_add_tcPr().append(sh)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
            if i in mono_cols: r.font.name = "Consolas"; r.font.size = Pt(size - 1)
    if widths:
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[i] * 1440)))
        for i, col in enumerate(t.columns): col.width = Inches(widths[i])
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs: p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.0
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def caption(d, text):
    p = d.add_paragraph(); r = p.add_run(text); r.italic = True; r.font.size = Pt(10.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; return p

def code(d, text, size=9.5):
    t = d.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.rows[0].cells[0]; c.text = ""
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "F4F4F4"); c._tc.get_or_add_tcPr().append(sh)
    lines = text.rstrip("\n").split("\n")
    p = c.paragraphs[0]; p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
    for i, ln in enumerate(lines):
        r = p.add_run(ln); r.font.name = "Consolas"; r.font.size = Pt(size)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if i < len(lines) - 1: r.add_break()
    d.add_paragraph().paragraph_format.space_after = Pt(2)

def cover(d, review_line, title, subtitle, extra_rows):
    for _ in range(2): d.add_paragraph()
    para(d, "VELLORE INSTITUTE OF TECHNOLOGY", bold=True, align="c", size=16, after=2)
    para(d, "School of Computer Science Engineering and Information Systems", align="c", size=12, after=18)
    para(d, "BCSE302L  Database Systems", bold=True, align="c", size=13, after=2)
    para(d, "Course Project", align="c", size=12, after=30)
    para(d, title, bold=True, align="c", size=22, after=4)
    para(d, subtitle, italic=True, align="c", size=13, after=30)
    para(d, review_line, bold=True, align="c", size=14, after=36)
    t = d.add_table(rows=0, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in [("Submitted by", "Vishwak"), ("Register Number", "24BYB0053")] + extra_rows + [("Faculty", "")]:
        cells = t.add_row().cells
        cells[0].text = ""; r = cells[0].paragraphs[0].add_run(k); r.bold = True
        cells[1].text = ""; cells[1].paragraphs[0].add_run(": " + v)
        cells[0].width = Inches(1.9); cells[1].width = Inches(3.4)
    d.add_paragraph(); d.add_paragraph()
    para(d, "October 2026", align="c", size=12)
    d.add_page_break()
